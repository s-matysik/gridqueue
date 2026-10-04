"""Adapter: Polskie Sieci Elektroenergetyczne S.A. (PL, operator PRZESYŁOWY).

Uwaga ustrojowa, którą trzeba zapisać w manuskrypcie, bo zmienia zakres wniosku:
PSE jest operatorem systemu PRZESYŁOWEGO, nie dystrybucyjnego.  Nie naciągamy tu
niczego -- dokument sam deklaruje podstawę: plik nazywa się "Informacje
publikowane zgodnie z Art. 7 ust. 8l ustawy PE", a komórka A1 arkusza powtarza tę
deklarację.  Art. 7 ust. 8l adresuje "przedsiębiorstwo energetyczne zajmujące się
przesyłaniem LUB dystrybucją", więc obowiązek obejmuje OSP tak samo jak OSD.

Co z tego naprawdę wynika dla porównywalności:
  * zakres napięć jest inny -- 110/220/400 kV zamiast SN/110 kV.  To nie jest
    niezgodność ze schematem (obowiązek dotyczy sieci powyżej 1 kV), tylko
    zmierzona różnica w polu ``poziom_napiecia``;
  * PSE realizuje pkt 1), 3) i 4) w jednym zestawieniu (encja WNIOSEK);
  * pkt 2) (dostępna moc w stacjach) PSE ujawnia INACZEJ: nie jako moc w MW, lecz
    jako "Listę rozdzielni, dla których brak jest dostępnych miejsc przyłączenia"
    z LICZBĄ dostępnych miejsc.  To jest informacja o węźle i o ograniczeniu, ale
    NIE jest wartością mocy -- mapujemy ją na ``ograniczenie_flaga`` i rozszerzenie
    ``_liczba_dostepnych_miejsc``, a pole ``moc_dostepna`` zostawiamy puste.
    Wpisanie tam zera albo liczby miejsc byłoby fabrykacją jednostki.

Nośnikiem jest XLSX, a nie PDF, więc geometria kolumn nie jest tu w ogóle problemem:
komórka arkusza ma tożsamość z definicji.  Zachowujemy jednak ten sam audyt
sprzeczności typu, żeby liczba dla PSE była porównywalna z liczbami dla publikujących
PDF-owych -- i żeby było widać, ile z niej pochodzi z układu, a ile z samego źródła.
"""

from __future__ import annotations

import re
from typing import Any, Optional

from ..schema import canonical_date
from . import _ptpiree as T
from .base import Adapter, ParseResult

__all__ = ["PSEAdapter"]

_SHEET_MAIN = "Wykaz wspólny"
_SHEET_VOCAB = "Listy rozwijane"


class PSEAdapter(Adapter):
    publisher = "PSE S.A."
    jurisdiction = "PL"
    declared_fields = T.DECLARED_FIELDS + ("id_wezla", "nazwa_wezla", "ograniczenie_flaga")
    declared_extensions = dict(
        T.DECLARED_EXTENSIONS,
        _liczba_dostepnych_miejsc="liczba wolnych miejsc przyłączenia w rozdzielni",
        _symbol_stacji="symbol stacji w nomenklaturze PSE",
        _zko_pse="Zakład Koordynacyjny Ośrodka PSE",
        _rozdzielnia="oznaczenie rozdzielni w stacji",
        _przypisy="odnośniki przypisów odcięte od treści komórek (pole=znacznik)",
    )
    detect_patterns = (r"pse", r"art_7_ust_8l", r"polskie sieci elektroenergetyczne")

    def __init__(self, data_publikacji: Optional[str] = None):
        self.data_publikacji = data_publikacji

    # ------------------------------------------------------------------
    def parse(self, source) -> ParseResult:
        """``source``: ścieżka do arkusza pkt 1/3/4 albo lista ścieżek.

        Rozpoznanie roli pliku po zawartości, nie po nazwie: arkusz o nazwie
        "Wykaz wspólny" to zestawienie wniosków, arkusz "Lista zajętych rozdzielni"
        to ujawnienie węzłowe.
        """
        import pandas as pd

        sources = [source] if isinstance(source, (str, bytes)) else list(source)
        frames, reports = [], []
        for src in sources:
            xl = pd.ExcelFile(src)
            if _SHEET_MAIN in xl.sheet_names:
                f, r = self._parse_wnioski(xl, src)
            elif any("rozdzielni" in s.lower() for s in xl.sheet_names):
                f, r = self._parse_wezly(xl, src)
            else:
                raise ValueError(f"nierozpoznany arkusz PSE w {src!r}: {xl.sheet_names}")
            frames.append(f)
            reports.append(r)
        df = pd.concat(frames, ignore_index=True) if frames else self.blank_frame(0)
        report = {
            "publikujacy": self.publisher,
            "strategia_wydobycia": "arkusz kalkulacyjny (tożsamość komórki z definicji)",
            "wzorzec": "PTPiREE/PSE, art. 7 ust. 8l pkt 1), 3), 4)",
            "dokumenty": reports,
            "n_wierszy": int(len(df)),
            "data_publikacji": self.data_publikacji,
        }
        return ParseResult(self.finalize(df, document="; ".join(map(str, sources))), report)

    # ------------------------------------------------------------------
    def _parse_wnioski(self, xl, src):
        import pandas as pd

        d, przypisy = _read_sheet_bez_przypisow(src, _SHEET_MAIN)
        tytul = " ".join(str(v) for v in d.iloc[0].tolist() if isinstance(v, str))
        stan = T.stan_na(tytul)
        if stan and self.data_publikacji is None:
            self.data_publikacji = stan

        hdr_row = self._find_header_row(d)
        body = d.iloc[hdr_row + 1:].reset_index(drop=True)
        body.columns = range(body.shape[1])

        # Ostatnie wiersze arkusza to LEGENDA przedrostków identyfikatora
        # ("E- Wnioskodawca wystąpił z wnioskiem...").  Wiersz danych ma wypełnione
        # pole wzorca "Nazwa OSD/OSP"; legenda go nie ma.  Kryterium jest więc
        # własnością wzorca, a nie zgadywaniem po pozycji.
        osd_col = 2
        keep = body[osd_col].notna() & body[osd_col].map(
        lambda v: bool(str(v).strip()) if v is not None else False)
        legenda = int((~keep).sum())
        body = body[keep].reset_index(drop=True)

        raw = pd.DataFrame(index=body.index)
        for phys, logical in T.PSE_XLSX_LAYOUT.items():
            if phys < body.shape[1]:
                raw[logical] = body[phys]
        out = T.to_schema(raw, publisher_key="pse", data_publikacji=self.data_publikacji)

        # Odnośniki przypisów wracają jako rozszerzenie, żeby odcięcie ich od
        # treści nie było utratą danej.  Klucz ``marks`` to (wiersz arkusza,
        # kolumna arkusza), a ``keep`` mówi, które wiersze arkusza zostały.
        wiersze_arkusza = [i for i, k in enumerate(keep.tolist(), start=hdr_row + 1) if k]
        znaczniki = []
        for ri in wiersze_arkusza:
            m = {T.PSE_XLSX_LAYOUT.get(ci, f"kol{ci}"): txt
                 for (rr, ci), txt in przypisy.items() if rr == ri}
            znaczniki.append("; ".join(f"{k}={v}" for k, v in sorted(m.items())) or None)
        out["_przypisy"] = znaczniki
        n_przypisow = sum(1 for z in znaczniki if z)

        vocab = self._read_vocab(xl)
        return out, {
            "plik": str(src),
            "arkusz": _SHEET_MAIN,
            "encja": "WNIOSEK",
            "n_kolumn_zrodlowych": int(body.shape[1]),
            "wiersz_naglowka": hdr_row,
            "wierszy_legendy_odrzuconych": legenda,
            "n_wierszy": int(len(out)),
            "stan_na": stan,
            "wierszy_z_przypisem": n_przypisow,
            "slownik_zrodlowy": vocab,
        }

    # ------------------------------------------------------------------
    def _parse_wezly(self, xl, src):
        """Lista rozdzielni bez wolnych miejsc przyłączenia (ujawnienie węzłowe)."""
        import pandas as pd

        sheet = next(s for s in xl.sheet_names if "rozdzielni" in s.lower())
        d = xl.parse(sheet, header=None)
        stan = T.stan_na(" ".join(str(v) for v in d.values.ravel() if isinstance(v, str)))

        hdr = None
        for i in range(min(12, len(d))):
            vals = [str(v).strip().lower() for v in d.iloc[i].tolist() if isinstance(v, str)]
            if any("symbol stacji" in v for v in vals):
                hdr = i
                break
        if hdr is None:
            raise ValueError(f"nie znaleziono nagłówka listy rozdzielni w {src!r}")
        cols = [str(v).strip() if isinstance(v, str) else "" for v in d.iloc[hdr].tolist()]
        body = d.iloc[hdr + 1:].reset_index(drop=True)
        body.columns = cols

        def c(frag):
            for name in body.columns:
                if frag in name.lower():
                    return body[name]
            return pd.Series([None] * len(body))

        symbol = c("symbol stacji").map(lambda v: None if v is None or (isinstance(v, float) and v != v) else str(v).strip())
        nazwa = c("nazwa stacji").map(lambda v: None if v is None or (isinstance(v, float) and v != v) else str(v).strip())
        rozdz = c("rozdzielnia").map(lambda v: None if v is None or (isinstance(v, float) and v != v) else str(v).strip())
        miejsc = c("liczba dostępnych")
        zko = c("zko")

        keep = symbol.notna() & nazwa.notna()
        idx = [i for i in range(len(body)) if keep.iloc[i]]

        out = pd.DataFrame(index=range(len(idx)))
        out["id_wniosku"] = None
        out["id_wezla"] = [f"pse:{symbol.iloc[i]}:{_volt_tag(rozdz.iloc[i])}" for i in idx]
        out["nazwa_wezla"] = [nazwa.iloc[i] for i in idx]
        out["lokalizacja_tekst"] = [nazwa.iloc[i] for i in idx]
        out["poziom_napiecia"] = [_volt_of(rozdz.iloc[i]) for i in idx]
        out["klasa_zasobu"] = "NIEOKRESLONA"
        out["status_procesu"] = "NIEOKRESLONY"
        # moc_dostepna NIE jest wypełniana: źródło podaje LICZBĘ MIEJSC, nie moc.
        out["moc_dostepna"] = None
        out["ograniczenie_flaga"] = [
            (float(miejsc.iloc[i]) == 0.0) if _num(miejsc.iloc[i]) is not None else None
            for i in idx
        ]
        out["data_publikacji"] = stan
        out["_encja"] = "WEZEL"
        out["_liczba_dostepnych_miejsc"] = [_num(miejsc.iloc[i]) for i in idx]
        out["_symbol_stacji"] = [symbol.iloc[i] for i in idx]
        out["_rozdzielnia"] = [rozdz.iloc[i] for i in idx]
        out["_zko_pse"] = [zko.iloc[i] for i in idx]
        return out, {
            "plik": str(src),
            "arkusz": sheet,
            "encja": "WEZEL",
            "n_wierszy": int(len(out)),
            "stan_na": stan,
            "uwaga": "źródło podaje liczbę wolnych miejsc przyłączenia, nie moc w MW; "
                     "moc_dostepna pozostaje pusta",
        }

    # ------------------------------------------------------------------
    @staticmethod
    def _find_header_row(d) -> int:
        for i in range(min(12, len(d))):
            vals = [str(v).strip().lower() for v in d.iloc[i].tolist() if isinstance(v, str)]
            if any(v.startswith("id obiektu") for v in vals):
                return i
        raise ValueError("nie znaleziono wiersza nagłówka arkusza PSE")

    @staticmethod
    def _read_vocab(xl) -> dict[str, list[str]]:
        """Słowniki kontrolowane zadeklarowane przez sam dokument."""
        if _SHEET_VOCAB not in xl.sheet_names:
            return {}
        lr = xl.parse(_SHEET_VOCAB, header=None)
        out: dict[str, list[str]] = {}
        for j in range(lr.shape[1]):
            vals = [re.sub(r"\s+", " ", str(v)).strip()
                    for v in lr[j].tolist() if isinstance(v, str) and str(v).strip()]
            if len(vals) >= 2:
                out[vals[0]] = vals[1:]
        return out


def _read_sheet_bez_przypisow(src, sheet: str):
    """Wczytaj arkusz, odłączając ODNOŚNIKI PRZYPISÓW od treści komórek.

    Dlaczego to nie jest kosmetyka.  Pomiar dokładności na wzorcu odczytanym
    wzrokowo pokazał siedem niezgodności u PSE i wszystkie miały jedną przyczynę:
    PSE oznacza część wartości odnośnikiem przypisu zapisanym jako INDEKS GÓRNY
    wewnątrz komórki -- lokalizacja "WYP" z przypisem M, "CPK" z przypisem K,
    "SE EJ" z F, "Norki" z Y, oraz moce "300,1", "50,1" i "300" z przypisem E.
    Płaski odczyt sklejał je w "WYPM", "CPKK", "300,1E" -- w polu tekstowym
    powstawał zanieczyszczony ciąg, a w polu liczbowym wartość PRZEPADAŁA, bo
    "300,1E" nie jest liczbą.

    Rozstrzygające jest to, że odnośnik daje się rozpoznać STRUKTURALNIE:
    openpyxl wystawia komórkę jako ``CellRichText``, a fragment przypisu ma
    ``font.vertAlign == "superscript"``.  Odcinamy więc po FORMATOWANIU, a nie
    przez zgadywanie, że końcowa wielka litera "pewnie jest przypisem" -- takie
    zgadywanie okaleczyłoby prawdziwe nazwy stacji kończące się literą.
    Sam odnośnik nie jest wyrzucany: wraca w rozszerzeniu ``_przypisy``.
    """
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock

    wb = load_workbook(src, rich_text=True, data_only=True)
    ws = wb[sheet]
    rows, marks = [], {}
    for ri, row in enumerate(ws.iter_rows(values_only=False)):
        out = []
        for ci, cell in enumerate(row):
            v = cell.value
            if isinstance(v, CellRichText):
                keep, sup = [], []
                for part in v:
                    if isinstance(part, TextBlock):
                        if getattr(part.font, "vertAlign", None) == "superscript":
                            sup.append(str(part.text))
                            continue
                        keep.append(str(part.text))
                    else:
                        keep.append(str(part))
                if sup:
                    marks[(ri, ci)] = "".join(sup)
                v = "".join(keep)
            out.append(v)
        rows.append(out)
    wb.close()
    return pd.DataFrame(rows), marks


def _num(v) -> Optional[float]:
    if v is None or (isinstance(v, float) and v != v):
        return None
    try:
        return float(str(v).strip().replace(",", "."))
    except ValueError:
        return None


def _volt_of(rozdzielnia: Any) -> Optional[str]:
    m = re.search(r"(\d+)\s*kV", str(rozdzielnia or ""), re.IGNORECASE)
    return f"{m.group(1)} kV" if m else None


def _volt_tag(rozdzielnia: Any) -> str:
    m = re.search(r"(\d+)", str(rozdzielnia or ""))
    return m.group(1) if m else "nn"
