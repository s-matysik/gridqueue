"""Adapter: ESB Networks (IE), network connection capacity map.

The publisher issues a single sheet covering **the entire network down to MV/LV
station level**, which makes it the most numerous source in the panel and also the
only one that goes below high-voltage station level. The row unit is a group of
transformers at a station.

Two units in one source
------------------------------
The publisher gives the **demand side in MVA** (apparent power), and the
**generation side in MW** (active power). The canonical unit of the schema is kW
of active power, so:

* generation values are converted (MW -> kW, the same physical quantity),
* demand values are **not converted** and go into extension fields
  with the unit explicitly recorded.

The `moc_dostepna` field therefore carries the available **firm generation**
capacity, not the sum of both sides. Summing MW with MVA would give a number with
no physical meaning, and choosing one side without saying so explicitly would be
hiding the decision.

The source gives coordinates and two separate constraint flags, for the demand
side and the generation side. The `ograniczenie_flaga` field receives the
**disjunction of both**, because the schema has a single field; both source values
remain in the extensions.
"""

from __future__ import annotations

import re
from pathlib import Path

from .base import Adapter, ParseResult

PUBLISHER = "ESB Networks DAC"
JURISDICTION = "IE"
SHEET = "Heatmap data"

#: The header spans two rows; the proper names are in the second one.
HEADER_ROW = 2

KOL = {
    "station": "Station Name",
    "group": "Transformer GroupID",
    "primary_voltage": "Primary kV",
    "secondary_voltage": "Secondary voltage(s)",
    "voltage_class": "Voltage Class",
    "configuration": "Transformer Configuration",
    "capacity_installed_mva": "Installed Capacity MVA",
    "withdrawal_permanent_mva": "Demand FirmCapacity MVA",
    "withdrawal_available_mva": "Demand Available MVA",
    "withdrawal_parent_mva": "Parent Available MVA",
    "withdrawal_limit": "Demand Parent Constraint",
    "gen_firm_mw": "Generation Firm Capacity MW",
    "gen_nonfirm_mw": "Generation NonFirm Capacity MW",
    "gen_contracted_mw": "Generation Total Committed MW",
    "gen_available_firm_mw": "Gen Available Firm MW",
    "gen_available_nonfirm_mw": "Gen Available NonFirm MW",
    "gen_restriction": "Generation Parent Constraint",
    "power_supply": "Parent Feeder",
    "parent_station": "Parent Station",
    "transmission_tie": "TSO Interface Station",
    "comment": "Comment",
    "latitude": "Latitude",
    "longitude": "Longitude",
}

_YES = {"yes", "y", "true", "constrained", "1"}
_NO = {"no", "n", "false", "unconstrained", "0"}

#: The constraint columns carry the phrase "Constrained, otherwise capacity available = X kVA"
#: or are empty. We assume that an empty cell means NO constraint, not an absence of a
#: declaration: the column is a binary indicator, where a positive value carries an
#: explanation, and a negative one is expressed by emptiness. This is an INTERPRETATION, not
#: content of the source (the publisher nowhere describes this convention), so the raw
#: phrase remains in the extension, and the decision is recorded in the adapter report.
_RESTRICTION_PHRASE = re.compile(r"constrained", re.IGNORECASE)
_CAPACITY_IN_PHRASE = re.compile(r"=\s*([\d.,]+)\s*kVA", re.IGNORECASE)


def _count(v) -> float | None:
    if v is None:
        return None
    s = str(v).strip().replace(",", "")
    if s in {"", "-", "_", "n/a", "N/A", "None", "nan"}:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _mw_per_kw(v) -> float | None:
    x = _count(v)
    return None if x is None else x * 1000.0


def _flag(v) -> bool | None:
    """Constraint indicator; an empty cell means no constraint (see above)."""
    if v is None or str(v).strip() == "":
        return False
    s = str(v).strip().lower()
    if s in _YES:
        return True
    if s in _NO:
        return False
    return True if _RESTRICTION_PHRASE.search(s) else None


def _unlimited_capacity(v) -> float | None:
    """Extract X from the phrase 'Constrained, otherwise capacity available = X kVA'.

    The value is in kVA, i.e. apparent power, so it is NOT converted to kW
    and stays in the extension with the unit in the field name.
    """
    if v is None:
        return None
    m = _CAPACITY_IN_PHRASE.search(str(v))
    return _count(m.group(1)) if m else None


def _alternative(a: bool | None, b: bool | None) -> bool | None:
    """Disjunction preserving indeterminacy: None or False -> None, not False."""
    if a is True or b is True:
        return True
    if a is False and b is False:
        return False
    return None


class ESBNetworksAdapter(Adapter):
    publisher = PUBLISHER
    jurisdiction = JURISDICTION
    declared_fields = (
        "id_wezla", "nazwa_wezla", "wspolrzedne", "poziom_napiecia",
        "moc_dostepna", "moc_zarezerwowana", "ograniczenie_flaga",
    )
    declared_extensions = {
        "_jednostka_mocy_odbiorczej": "unit of the demand side at the publisher (MVA)",
        "_klasa_napiecia": "voltage class per the publisher",
        "_konfiguracja_transformatorow": "configuration of the transformer group",
        "_moc_zainstalowana_mva": "installed capacity of the transformers",
        "_odbior_dostepna_mva": "available demand capacity",
        "_odbior_trwala_mva": "firm demand capacity",
        "_odbior_nadrzedna_mva": "available demand capacity of the parent station",
        "_odbior_ograniczenie": "constraint on the demand side",
        "_wytw_ograniczenie": "constraint on the generation side",
        "_odbior_ograniczenie_opis": "raw constraint phrase on the demand side",
        "_wytw_ograniczenie_opis": "raw constraint phrase on the generation side",
        "_odbior_bez_ograniczenia_kva": "demand capacity available if not for the constraint",
        "_wytw_bez_ograniczenia_kva": "generation capacity available if not for the constraint",
        "_wytw_dostepna_nietrwala_kw": "available non-firm generation capacity",
        "_wytw_zakontraktowana_kw": "total contracted generation capacity",
        "_wytw_trwala_kw": "firm generation capacity",
        "_zasilanie": "feeder line",
        "_stacja_nadrzedna": "parent station",
        "_styk_przesylowy": "station of interface with the transmission network",
        "_komentarz": "publisher's comment",
    }
    detect_patterns = (r"customer.?heatmap", r"esb.?networks", r"heatmap.?data")

    def parse(self, source) -> ParseResult:
        import openpyxl
        import pandas as pd

        path = Path(str(source))
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        for _ in range(HEADER_ROW - 1):
            next(it)
        header = [str(c).strip() if c is not None else "" for c in next(it)]
        position = {name: header.index(label)
                   for name, label in KOL.items() if label in header}
        missing = sorted(set(KOL) - set(position))

        records, without_coordinates, limited = [], 0, 0
        for w in it:
            def g(name):
                i = position.get(name)
                if i is None or i >= len(w):
                    return None
                v = w[i]
                return None if v is None or str(v).strip() in {"", "_", "-"} else v

            station = g("station")
            if station is None:
                continue
            group = g("group")
            lat, lon = _count(g("latitude")), _count(g("longitude"))
            if lat is None or lon is None:
                without_coordinates += 1
            s_withdrawal, s_generation = g("withdrawal_limit"), g("gen_restriction")
            f_withdrawal, f_generation = _flag(s_withdrawal), _flag(s_generation)
            limit = _alternative(f_withdrawal, f_generation)
            if limit:
                limited += 1
            voltage = g("primary_voltage")

            records.append({
                "id_wezla": f"{station}|{group}" if group else str(station),
                "nazwa_wezla": str(station).strip(),
                "wspolrzedne": f"{lat},{lon}" if lat is not None and lon is not None else None,
                "poziom_napiecia": str(voltage).strip() if voltage else None,
                # The generation side is in MW of active power, the only one that is convertible.
                "moc_dostepna": _mw_per_kw(g("gen_available_firm_mw")),
                "moc_zarezerwowana": _mw_per_kw(g("gen_contracted_mw")),
                "ograniczenie_flaga": limit,
                "_encja": "WEZEL",
                "_jednostka_mocy_odbiorczej": "MVA",
                "_klasa_napiecia": g("voltage_class"),
                "_konfiguracja_transformatorow": g("configuration"),
                "_moc_zainstalowana_mva": _count(g("capacity_installed_mva")),
                "_odbior_dostepna_mva": _count(g("withdrawal_available_mva")),
                "_odbior_trwala_mva": _count(g("withdrawal_permanent_mva")),
                "_odbior_nadrzedna_mva": _count(g("withdrawal_parent_mva")),
                "_odbior_ograniczenie": f_withdrawal,
                "_wytw_ograniczenie": f_generation,
                "_odbior_ograniczenie_opis": str(s_withdrawal) if s_withdrawal is not None else None,
                "_wytw_ograniczenie_opis": str(s_generation) if s_generation is not None else None,
                "_odbior_bez_ograniczenia_kva": _unlimited_capacity(s_withdrawal),
                "_wytw_bez_ograniczenia_kva": _unlimited_capacity(s_generation),
                "_wytw_dostepna_nietrwala_kw": _mw_per_kw(g("gen_available_nonfirm_mw")),
                "_wytw_zakontraktowana_kw": _mw_per_kw(g("gen_contracted_mw")),
                "_wytw_trwala_kw": _mw_per_kw(g("gen_firm_mw")),
                "_zasilanie": g("power_supply"),
                "_stacja_nadrzedna": g("parent_station"),
                "_styk_przesylowy": g("transmission_tie"),
                "_komentarz": g("comment"),
            })
        wb.close()

        df = self.finalize(pd.DataFrame(records), document=path.name)
        notes = {
            "publisher": PUBLISHER,
            "jurisdiction": JURISDICTION,
            "release_mode": "spreadsheet on the publisher's website",
            "rows": int(len(df)),
            "rows_without_coordinates": without_coordinates,
            "constrained_nodes": limited,
            "empty_flag_interpretation": (
                "empty cell of the constraint column read as NO constraint; "
                "the publisher does not describe this convention, raw phrase in the extension"
            ),
            "columns_not_found": missing,
            "generating_side_unit": "MW",
            "receiving_side_unit": "MVA",
            "note_units": (
                "The moc_dostepna field carries the available firm GENERATION capacity, "
                "converted from MW to kW. The demand side is given in MVA (apparent power) "
                "and is not converted nor summed with the generation side; it remains in "
                "the extensions."
            ),
        }
        return ParseResult(frame=df, report=notes)
