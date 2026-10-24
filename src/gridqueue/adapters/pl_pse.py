"""Adapter: Polskie Sieci Elektroenergetyczne S.A. (PL, TRANSMISSION operator).

Structural note that must be recorded in the manuscript, because it changes the scope of
the claim:
PSE is the operator of the TRANSMISSION system, not the distribution system. We are not
stretching anything here: the document itself declares the basis: the file is named "Informacje
publikowane zgodnie z Art. 7 ust. 8l ustawy PE", and cell A1 of the worksheet repeats that
declaration. Art. 7(8l) addresses an "energy undertaking engaged in transmission OR
distribution", so the duty covers the OSP just as much as the OSD.

What this really means for comparability:
  * the voltage range is different: 110/220/400 kV instead of MV/110 kV. This is not a
    mismatch with the schema (the duty applies to networks above 1 kV), just a measured
    difference in the ``voltage_level`` field;
  * PSE fulfills items 1), 3) and 4) in a single compilation (the APPLICATION entity);
  * item 2) (available capacity at the stations) PSE discloses DIFFERENTLY: not as capacity
    in MW, but as a "Listę rozdzielni, dla których brak jest dostępnych miejsc przyłączenia"
    with the NUMBER of available slots. This is information about the node and about the
    constraint, but it is NOT a capacity value: we map it to ``constraint_flag`` and the
    extension ``_free_slots``, and we leave the ``capacity_available`` field empty.
    Putting a zero or the number of slots there would be a fabrication of the unit.

The carrier is XLSX, not PDF, so column geometry is not a problem here at all: the worksheet
cell has identity by definition. We nonetheless keep the same type-mismatch audit, so that
the count for PSE is comparable to the counts for PDF-based publishers, and so that it is
visible how much of it comes from the layout and how much from the source itself.
"""

from __future__ import annotations

import re
from typing import Any, Optional

from ..schema import canonical_date
from . import _ptpiree as T
from .base import Adapter, ParseResult

__all__ = ["PSEAdapter"]

_SHEET_MAIN = "Wykaz wspólny"
_SHEET_VOCAB = "Listy rozwijane"


class PSEAdapter(Adapter):
    publisher = "PSE S.A."
    jurisdiction = "PL"
    declared_fields = T.DECLARED_FIELDS + ("node_id", "node_name", "constraint_flag")
    declared_extensions = dict(
        T.DECLARED_EXTENSIONS,
        _free_slots="number of free connection slots at the switching station",
        _station_symbol="station symbol in PSE nomenclature",
        _pse_coordination_centre="PSE Coordination Centre",
        _switchgear="identifier of the switching station within the station",
        _footnotes="footnote references stripped from the cell contents (field=marker)",
    )
    detect_patterns = (r"pse", r"art_7_ust_8l", r"polskie sieci elektroenergetyczne")

    def __init__(self, publication_date: Optional[str] = None):
        self.publication_date = publication_date

    # ------------------------------------------------------------------
    def parse(self, source) -> ParseResult:
        """``source``: path to the worksheet for items 1/3/4, or a list of paths.

        Recognizing the role of the file by its content, not by its name: a worksheet named
        "Wykaz wspólny" is the compilation of applications, the worksheet "Lista zajętych
        rozdzielni" is the node-level disclosure.
        """
        import pandas as pd

        sources = [source] if isinstance(source, (str, bytes)) else list(source)
        frames, reports = [], []
        for src in sources:
            xl = pd.ExcelFile(src)
            if _SHEET_MAIN in xl.sheet_names:
                f, r = self._parse_applications(xl, src)
            elif any("rozdzielni" in s.lower() for s in xl.sheet_names):
                f, r = self._parse_nodes(xl, src)
            else:
                raise ValueError(f"unrecognized PSE worksheet in {src!r}: {xl.sheet_names}")
            frames.append(f)
            reports.append(r)
        df = pd.concat(frames, ignore_index=True) if frames else self.blank_frame(0)
        report = {
            "publisher": self.publisher,
            "extraction_strategy": "spreadsheet (cell identity by definition)",
            "pattern": "PTPiREE/PSE, Art. 7(8l)(1), (3), (4)",
            "documents": reports,
            "n_rows": int(len(df)),
            "publication_date": self.publication_date,
        }
        return ParseResult(self.finalize(df, document="; ".join(map(str, sources))), report)

    # ------------------------------------------------------------------
    def _parse_applications(self, xl, src):
        import pandas as pd

        d, footnotes = _read_sheet_without_footnotes(src, _SHEET_MAIN)
        title = " ".join(str(v) for v in d.iloc[0].tolist() if isinstance(v, str))
        as_at = T.parse_as_at(title)
        if as_at and self.publication_date is None:
            self.publication_date = as_at

        hdr_row = self._find_header_row(d)
        body = d.iloc[hdr_row + 1:].reset_index(drop=True)
        body.columns = range(body.shape[1])

        # The last rows of the worksheet are the LEGEND of identifier prefixes
        # ("E- Wnioskodawca wystąpił z wnioskiem..."). A data row has the pattern
        # field "Nazwa OSD/OSP" filled in; the legend does not have it. The criterion is
        # thus a property of the pattern, not a guess based on position.
        osd_col = 2
        keep = body[osd_col].notna() & body[osd_col].map(
        lambda v: bool(str(v).strip()) if v is not None else False)
        legend = int((~keep).sum())
        body = body[keep].reset_index(drop=True)

        raw = pd.DataFrame(index=body.index)
        for phys, logical in T.PSE_XLSX_LAYOUT.items():
            if phys < body.shape[1]:
                raw[logical] = body[phys]
        out = T.to_schema(raw, publisher_key="pse", publication_date=self.publication_date)

        # Footnote references come back as an extension, so that cutting them off from
        # the content is not a loss of data. The ``marks`` key is (worksheet row,
        # worksheet column), and ``keep`` tells which worksheet rows remained.
        sheet_rows = [i for i, k in enumerate(keep.tolist(), start=hdr_row + 1) if k]
        markers = []
        for ri in sheet_rows:
            m = {T.PSE_XLSX_LAYOUT.get(ci, f"kol{ci}"): txt
                 for (rr, ci), txt in footnotes.items() if rr == ri}
            markers.append("; ".join(f"{k}={v}" for k, v in sorted(m.items())) or None)
        out["_footnotes"] = markers
        n_footnotes = sum(1 for z in markers if z)

        vocab = self._read_vocab(xl)
        return out, {
            "file_name": str(src),
            "sheet": _SHEET_MAIN,
            "entity": "APPLICATION",
            "n_source_columns": int(body.shape[1]),
            "header_row": hdr_row,
            "rejected_legend_rows": legend,
            "n_rows": int(len(out)),
            "parse_as_at": as_at,
            "rows_with_footnote": n_footnotes,
            "source_dictionary": vocab,
        }

    # ------------------------------------------------------------------
    def _parse_nodes(self, xl, src):
        """List of switching stations without free connection slots (node-level disclosure)."""
        import pandas as pd

        sheet = next(s for s in xl.sheet_names if "rozdzielni" in s.lower())
        d = xl.parse(sheet, header=None)
        as_at = T.parse_as_at(" ".join(str(v) for v in d.values.ravel() if isinstance(v, str)))

        hdr = None
        for i in range(min(12, len(d))):
            vals = [str(v).strip().lower() for v in d.iloc[i].tolist() if isinstance(v, str)]
            if any("symbol stacji" in v for v in vals):
                hdr = i
                break
        if hdr is None:
            raise ValueError(f"header of the switching-station list not found in {src!r}")
        cols = [str(v).strip() if isinstance(v, str) else "" for v in d.iloc[hdr].tolist()]
        body = d.iloc[hdr + 1:].reset_index(drop=True)
        body.columns = cols

        def c(frag):
            for name in body.columns:
                if frag in name.lower():
                    return body[name]
            return pd.Series([None] * len(body))

        symbol = c("symbol stacji").map(lambda v: None if v is None or (isinstance(v, float) and v != v) else str(v).strip())
        label_name = c("nazwa stacji").map(lambda v: None if v is None or (isinstance(v, float) and v != v) else str(v).strip())
        switchgear = c("rozdzielnia").map(lambda v: None if v is None or (isinstance(v, float) and v != v) else str(v).strip())
        places = c("liczba dostępnych")
        zko = c("zko")

        keep = symbol.notna() & label_name.notna()
        idx = [i for i in range(len(body)) if keep.iloc[i]]

        out = pd.DataFrame(index=range(len(idx)))
        out["application_id"] = None
        out["node_id"] = [f"pse:{symbol.iloc[i]}:{_volt_tag(switchgear.iloc[i])}" for i in idx]
        out["node_name"] = [label_name.iloc[i] for i in idx]
        out["location_text"] = [label_name.iloc[i] for i in idx]
        out["voltage_level"] = [_volt_of(switchgear.iloc[i]) for i in idx]
        out["resource_class"] = "UNSPECIFIED"
        out["process_status"] = "UNSPECIFIED"
        # capacity_available is NOT filled in: the source gives the NUMBER OF SLOTS, not capacity.
        out["capacity_available"] = None
        out["constraint_flag"] = [
            (float(places.iloc[i]) == 0.0) if _num(places.iloc[i]) is not None else None
            for i in idx
        ]
        out["date_published"] = as_at
        out["_entity"] = "NODE"
        out["_free_slots"] = [_num(places.iloc[i]) for i in idx]
        out["_station_symbol"] = [symbol.iloc[i] for i in idx]
        out["_switchgear"] = [switchgear.iloc[i] for i in idx]
        out["_pse_coordination_centre"] = [zko.iloc[i] for i in idx]
        return out, {
            "file_name": str(src),
            "sheet": sheet,
            "entity": "NODE",
            "n_rows": int(len(out)),
            "parse_as_at": as_at,
            "note": "the source gives the number of free connection slots, not capacity in "
                     "MW; capacity_available remains empty",
        }

    # ------------------------------------------------------------------
    @staticmethod
    def _find_header_row(d) -> int:
        for i in range(min(12, len(d))):
            vals = [str(v).strip().lower() for v in d.iloc[i].tolist() if isinstance(v, str)]
            if any(v.startswith("id obiektu") for v in vals):
                return i
        raise ValueError("header row of the PSE worksheet not found")

    @staticmethod
    def _read_vocab(xl) -> dict[str, list[str]]:
        """Controlled vocabularies declared by the document itself."""
        if _SHEET_VOCAB not in xl.sheet_names:
            return {}
        lr = xl.parse(_SHEET_VOCAB, header=None)
        out: dict[str, list[str]] = {}
        for j in range(lr.shape[1]):
            vals = [re.sub(r"\s+", " ", str(v)).strip()
                    for v in lr[j].tolist() if isinstance(v, str) and str(v).strip()]
            if len(vals) >= 2:
                out[vals[0]] = vals[1:]
        return out


def _read_sheet_without_footnotes(src, sheet: str):
    """Read the worksheet, detaching FOOTNOTE REFERENCES from the cell contents.

    Why this is not cosmetic. An accuracy measurement against a reference set read
    visually showed seven mismatches for PSE, and they all had a single cause:
    PSE marks part of the values with a footnote reference written as a SUPERSCRIPT
    inside the cell: the location "WYP" with footnote M, "CPK" with footnote K,
    "SE EJ" with F, "Norki" with Y, and the capacities "300,1", "50,1" and "300" with
    footnote E. A flat read glued them into "WYPM", "CPKK", "300,1E": in the text field
    a contaminated string was produced, and in the numeric field the value WAS LOST,
    because "300,1E" is not a number.

    What is decisive is that the reference can be recognized STRUCTURALLY:
    openpyxl exposes the cell as ``CellRichText``, and the footnote fragment has
    ``font.vertAlign == "superscript"``. So we cut it off based on FORMATTING, not
    by guessing that a trailing capital letter "is probably a footnote": such
    guessing would mutilate real station names ending in a letter.
    The reference itself is not discarded: it comes back in the ``_footnotes`` extension.
    """
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock

    wb = load_workbook(src, rich_text=True, data_only=True)
    ws = wb[sheet]
    rows, marks = [], {}
    for ri, row in enumerate(ws.iter_rows(values_only=False)):
        out = []
        for ci, cell in enumerate(row):
            v = cell.value
            if isinstance(v, CellRichText):
                keep, sup = [], []
                for part in v:
                    if isinstance(part, TextBlock):
                        if getattr(part.font, "vertAlign", None) == "superscript":
                            sup.append(str(part.text))
                            continue
                        keep.append(str(part.text))
                    else:
                        keep.append(str(part))
                if sup:
                    marks[(ri, ci)] = "".join(sup)
                v = "".join(keep)
            out.append(v)
        rows.append(out)
    wb.close()
    return pd.DataFrame(rows), marks


def _num(v) -> Optional[float]:
    if v is None or (isinstance(v, float) and v != v):
        return None
    try:
        return float(str(v).strip().replace(",", "."))
    except ValueError:
        return None


def _volt_of(switchgear: Any) -> Optional[str]:
    m = re.search(r"(\d+)\s*kV", str(switchgear or ""), re.IGNORECASE)
    return f"{m.group(1)} kV" if m else None


def _volt_tag(switchgear: Any) -> str:
    m = re.search(r"(\d+)", str(switchgear or ""))
    return m.group(1) if m else "nn"
